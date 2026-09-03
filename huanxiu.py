import os
import datetime
import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import load_workbook
from demos import Crili
import config


class App:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("串休管理")
        self.root.geometry("480x280")
        self.root.resizable(False, False)
        self.root.eval("tk::PlaceWindow . center")

        style = ttk.Style()
        style.theme_use("vista")
        style.configure("Display.TEntry", foreground="#555")

        font = ("Microsoft YaHei", 14)
        self.root.option_add("*Font", font)

        main_frame = ttk.Frame(self.root, padding="30 25 30 12")
        main_frame.pack(fill=tk.BOTH, expand=True)

        main_frame.columnconfigure(1, weight=1)
        main_frame.columnconfigure(3, weight=1)

        self.id_name_map = {}
        self.load_data()

        row_h = 14

        ttk.Label(main_frame, text="职号：").grid(
            row=0, column=0, sticky=tk.E, pady=row_h)
        self.id_var = tk.StringVar()
        self.entry_id = ttk.Entry(main_frame, textvariable=self.id_var)
        self.entry_id.grid(row=0, column=1, sticky=tk.EW,
                           pady=row_h, padx=(6, 0))

        ttk.Label(main_frame, text="姓名：").grid(
            row=0, column=2, sticky=tk.E, pady=row_h, padx=(20, 0))
        self.name_var = tk.StringVar()
        self.entry_name = ttk.Entry(
            main_frame, state="readonly", textvariable=self.name_var)
        self.entry_name.grid(row=0, column=3, sticky=tk.EW,
                             pady=row_h, padx=(6, 0))

        # 年月输入框（默认显示全局配置的年月，格式 YYYY-MM，用户仍可在界面手动修改）
        ttk.Label(main_frame, text="年月：").grid(
            row=1, column=0, sticky=tk.E, pady=row_h)
        self.month_var = tk.StringVar()
        default_ym = f"{config.YEAR}-{config.MONTH:02d}"
        self.month_var.set(default_ym)
        self.entry_month = ttk.Entry(main_frame, textvariable=self.month_var)
        self.entry_month.grid(
            row=1, column=1, sticky=tk.EW, pady=row_h, padx=(6, 0))

        ttk.Label(main_frame, text="串休时长：").grid(
            row=2, column=0, sticky=tk.E, pady=row_h)
        self.entry_hours = ttk.Entry(main_frame)
        self.entry_hours.grid(row=2, column=1, columnspan=3,
                              sticky=tk.EW, pady=row_h, padx=(6, 0))

        btn_frame = ttk.Frame(main_frame)
        btn_frame.grid(row=3, column=0, columnspan=4, pady=(18, 4))
        self.btn_submit = ttk.Button(
            btn_frame, text="抵扣串休", style="Big.TButton", command=self.on_deduct_hours)
        self.btn_clear = ttk.Button(
            btn_frame, text="清空 K 列", style="Big.TButton", command=self.on_clear_k_column)

        style.configure("Big.TButton", font=("Microsoft YaHei", 14))
        self.btn_submit.pack(side=tk.LEFT, padx=(0, 8))
        self.btn_clear.pack(side=tk.LEFT)

        self.id_var.trace_add("write", self.on_id_changed)
        self.entry_id.focus()
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    def load_data(self):
        excel_path = os.path.join(os.path.dirname(__file__), "计算结果.xlsx")
        if not os.path.exists(excel_path):
            messagebox.showerror("错误", f"找不到文件：{excel_path}")
            return
        try:
            wb = load_workbook(excel_path, read_only=True)
            if "记录表" not in wb.sheetnames:
                messagebox.showerror("错误", '找不到工作表 "记录表"')
                wb.close()
                return
            ws = wb["记录表"]
            count = 0
            for row in ws.iter_rows(min_row=3, min_col=1, max_col=2, values_only=True):
                if row[1] is not None:
                    key = str(int(row[1])) if isinstance(
                        row[1], float) else str(row[1]).strip()
                    val = str(row[0]).strip() if row[0] is not None else ""
                    self.id_name_map[key] = val
                    count += 1
            wb.close()
            self.status_text = f"已加载 {count} 条记录"
        except Exception as e:
            messagebox.showerror("错误", f"读取 Excel 失败：{e}")

    def on_id_changed(self, *_):
        id_ = self.id_var.get().strip()
        if not id_:
            self.name_var.set("")
            return
        name = self.id_name_map.get(id_, "")
        if name:
            self.name_var.set(name)
        else:
            self.name_var.set("")

    def on_deduct_hours(self):
        id_ = self.id_var.get().strip()
        ym_text = self.month_var.get().strip()
        duration_text = self.entry_hours.get().strip()

        if not id_:
            messagebox.showerror("错误", "请先输入职号。")
            return
        if not ym_text:
            messagebox.showerror("错误", "请先输入年月。")
            return
        if not duration_text:
            messagebox.showerror("错误", "请先输入串休时长。")
            return

        try:
            year_str, month_str = ym_text.split("-")
            year = int(year_str)
            month = int(month_str)
            if not (1 <= month <= 12):
                raise ValueError
        except Exception:
            messagebox.showerror("错误", "年月格式应为 YYYY-MM，例如 2026-05。")
            return

        try:
            remain_hours = float(duration_text)
        except ValueError:
            messagebox.showerror("错误", "串休时长必须是数字。")
            return

        try:
            crili = Crili(year, month)
            result = crili.parseHTML()
        except Exception as e:
            messagebox.showerror("错误", f"获取加班倍率失败：{e}")
            return

        excel_path = os.path.join(os.path.dirname(__file__), "计算结果.xlsx")
        if not os.path.exists(excel_path):
            messagebox.showerror("错误", f"找不到文件：{excel_path}")
            return

        try:
            wb = load_workbook(excel_path)
            ws = wb["统计表"]

            def to_date(val):
                if isinstance(val, datetime.datetime):
                    return val
                if isinstance(val, str):
                    for fmt in ["%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"]:
                        try:
                            return datetime.datetime.strptime(val, fmt)
                        except ValueError:
                            continue
                return None

            rows = list(ws.iter_rows(min_row=2))
            # 先收集所有加班记录（包含日期、倍率、总加班时长）
            all_overtime_records = []
            for row in rows:
                rid = row[1].value
                if rid is None:
                    continue
                rid = str(rid).strip()
                if rid != id_:
                    continue
                day = to_date(row[2].value)
                if day is None:
                    continue
                if day.year != year or day.month != month:
                    continue
                multiplier = result.get(day.strftime("%Y%m%d"), 1.5)
                overtime_hours = row[8].value if row[8].value is not None else 0
                try:
                    overtime_hours = float(overtime_hours)
                except Exception:
                    overtime_hours = 0
                if overtime_hours > 0:
                    all_overtime_records.append((multiplier, row, day, overtime_hours))
            
            if not all_overtime_records:
                messagebox.showwarning("警告", "该人员当月无加班记录。")
                return
            
            # 按日期分组，计算每组的总加班时长和已抵扣串休
            from collections import defaultdict
            day_data = defaultdict(lambda: {"total_hours": 0, "deducted": 0, "multiplier": 1.5, "row": None})
            for multiplier, row, day, overtime_hours in all_overtime_records:
                day_key = day.strftime("%Y%m%d")
                day_data[day_key]["total_hours"] += overtime_hours
                day_data[day_key]["multiplier"] = max(multiplier, day_data[day_key]["multiplier"])
                day_data[day_key]["row"] = row
            
            # 计算每组的已抵扣串休
            for day_key, data in day_data.items():
                row = data["row"]
                k_value = row[10].value or 0
                try:
                    k_value = float(k_value)
                except:
                    k_value = 0
                data["deducted"] = k_value  # k_value 是负数
            
            # 构建用于抵扣的列表
            overtime_rows = []
            for day_key, data in day_data.items():
                # 剩余可用加班时长 = 总加班时长 - |已抵扣串休|
                remaining = data["total_hours"] - (-data["deducted"])
                if remaining <= 0:
                    continue
                overtime_rows.append((data["multiplier"], data["row"], remaining))

            # 计算剩余可用加班时长
            total_overtime = sum(item[2] for item in overtime_rows)
            remaining_overtime = total_overtime

            if remaining_overtime <= 0:
                messagebox.showwarning("警告", "该人员当月加班时长已全部抵扣完毕。")
                return

            # 按工作日、公休日、节假日顺序抵扣
            overtime_rows.sort(key=lambda item: (
                item[0] != 1.5, item[0], item[0] != 2))

            used_total = 0.0
            for multiplier, row, avail in overtime_rows:
                if remain_hours <= 0:
                    break
                used = min(avail, remain_hours)
                # 累加 K 列已有的抵扣值，而不是覆盖
                existing_k = row[10].value or 0
                row[10].value = existing_k - used
                remain_hours -= used
                used_total += used

            if remain_hours > 0:
                wb.close()
                messagebox.showerror(
                    "错误", f"加班时长不足，剩余可用 {remaining_overtime:.2f} 小时，"
                           f"仍需抵扣 {remain_hours:.2f} 小时。")
                return

            # 计算累计抵扣值（按日期分组后计算，避免重复）
            total_deducted_after = 0.0
            for day_key, data in day_data.items():
                row = data["row"]
                k_value = row[10].value or 0
                try:
                    k_value = float(k_value)
                    total_deducted_after += k_value
                except:
                    pass

            wb.save(excel_path)
            wb.close()
            self.entry_hours.delete(0, tk.END)  # 运行后清空输入框
            messagebox.showinfo(
                "成功", f"本次抵扣 {used_total:.2f} 小时，累计已抵扣 {abs(total_deducted_after):.2f} 小时。\n请检查统计表 K 列。")
        except Exception as e:
            messagebox.showerror("错误", f"处理 Excel 失败：{e}")
            return

    def on_clear_k_column(self):
        id_ = self.id_var.get().strip()
        ym_text = self.month_var.get().strip()

        if not id_:
            messagebox.showerror("错误", "请先输入职号。")
            return
        if not ym_text:
            messagebox.showerror("错误", "请先输入年月。")
            return

        try:
            year_str, month_str = ym_text.split("-")
            year = int(year_str)
            month = int(month_str)
            if not (1 <= month <= 12):
                raise ValueError
        except Exception:
            messagebox.showerror("错误", "年月格式应为 YYYY-MM，例如 2026-05。")
            return

        excel_path = os.path.join(os.path.dirname(__file__), "计算结果.xlsx")
        if not os.path.exists(excel_path):
            messagebox.showerror("错误", f"找不到文件：{excel_path}")
            return

        try:
            wb = load_workbook(excel_path)
            ws = wb["统计表"]

            def to_date(val):
                if isinstance(val, datetime.datetime):
                    return val
                if isinstance(val, str):
                    for fmt in ["%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"]:
                        try:
                            return datetime.datetime.strptime(val, fmt)
                        except ValueError:
                            continue
                return None

            cleared = 0
            for row in ws.iter_rows(min_row=2):
                rid = row[1].value
                if rid is None:
                    continue
                rid = str(rid).strip()
                if rid != id_:
                    continue
                day = to_date(row[2].value)
                if day is None:
                    continue
                if day.year != year or day.month != month:
                    continue
                if row[10].value is not None:
                    row[10].value = None
                    cleared += 1

            wb.save(excel_path)
            wb.close()
            messagebox.showinfo("成功", f"已清空 {cleared} 个 K 列值。")
        except Exception as e:
            messagebox.showerror("错误", f"清空 K 列失败：{e}")
            return

    def on_close(self):
        self.root.destroy()

    def run(self):
        self.root.mainloop()


if __name__ == "__main__":
    App().run()
