from demos import *
import sys
import logging
from openpyxl import load_workbook
from process_attendance import calculate_main

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

if __name__ == "__main__":
    cw = Cwindow()
    if not cw.createWindow():
        logging.info("已取消计算：窗口已关闭")
        sys.exit(0)

    night_num = ['Q4642', '60836']  # 60836
    # 获得工作日和节假日
    start = time.perf_counter()
    wb = load_workbook(filename="计算结果.xlsx")
    ws = wb["记录表"]
    # 用id取代name，识别职号
    ids_criterion = []
    ids_night = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=2, max_col=2):
        for id in row:
            if id.value is not None and id.value not in night_num:
                ids_criterion.append(id.value)
            elif id.value in night_num:
                ids_night.append(id.value)
    
    # 统一使用 calculate_main 计算，夜班人员通过 night_ids 参数传入
    all_ids = ids_criterion + ids_night
    logging.info(f"开始计算加班，标准人员: {len(ids_criterion)} 人, 夜班人员: {len(ids_night)} 人")
    calculate_main(all_ids, night_ids=ids_night)

    logging.info(f"运行时间：{time.perf_counter() - start:.2f}秒")
